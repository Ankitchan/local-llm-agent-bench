"""Scoring logic for the three tasks. The check logic is unchanged from the version used for the published results;
only path handling differs. NOTE: candidate scripts are LLM-generated code. They are executed by score.py (see README warning)."""
import copy, importlib.util, re, sys
import openpyxl, pandas as pd, docx


def check_merge_intervals(solution_path):
    """Task A. Returns (cases_passed, cases_total, failure_messages)."""
    spec = importlib.util.spec_from_file_location("sol", solution_path)
    m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
    f = m.merge_intervals
    cases = [
        ([[1, 3], [2, 6], [8, 10], [15, 18]], [[1, 6], [8, 10], [15, 18]]),
        ([[8, 10], [1, 3], [2, 6]], [[1, 6], [8, 10]]),   # unsorted
        ([[1, 3], [3, 5]], [[1, 5]]),                      # touching
        ([[1, 10], [2, 3], [4, 5]], [[1, 10]]),            # nested
        ([], []),
        ([[5, 5]], [[5, 5]]),
    ]
    fails = []
    for inp, exp in cases:
        orig = copy.deepcopy(inp)
        got = [list(x) for x in f(inp)]
        if got != exp or inp != orig:
            fails.append(f"{orig} -> {got} expected {exp}" + (" (mutated input)" if inp != orig else ""))
    return len(cases) - len(fails), len(cases), fails


def check_xlsx(xlsx_path, csv_path):
    """Task B. Returns {check_name: bool}."""
    wb = openpyxl.load_workbook(xlsx_path)
    df = pd.read_csv(csv_path, parse_dates=["date"]); ex = df[df.amount < 0].copy(); ex["spent"] = -ex.amount
    exp_tot = ex.groupby("category").spent.sum().round(2); exp_cnt = ex.groupby("category").size()
    exp_mon = ex.assign(m=ex.date.dt.strftime("%Y-%m")).pivot_table(index="m", columns="category", values="spent", aggfunc="sum").fillna(0)
    checks = {"sheets": {"By Category", "Monthly"} <= set(wb.sheetnames)}
    def rows(ws): return [[c.value for c in r] for r in ws.iter_rows()]
    if checks["sheets"]:
        r = rows(wb["By Category"]); got = {}
        for row in r:
            if row and row[0] in exp_tot.index and isinstance(row[1], (int, float)):
                got[row[0]] = (row[1], row[2])
        checks["cat_totals"] = len(got) == len(exp_tot) and all(abs(got[k][0] - exp_tot[k]) < 0.011 for k in exp_tot.index if k in got)
        checks["cat_counts"] = len(got) == len(exp_tot) and all(got[k][1] == exp_cnt[k] for k in exp_cnt.index if k in got)
        order = list(exp_tot.sort_values(ascending=False).index)
        checks["sorted_desc"] = [row[0] for row in r if row and row[0] in exp_tot.index] == order
        tot = [row for row in r if row and isinstance(row[0], str) and row[0].strip().upper() == "TOTAL"]
        checks["total_formula"] = bool(tot) and any(isinstance(v, str) and v.upper().startswith("=SUM(") for v in tot[0][1:])
        m = rows(wb["Monthly"]); hdr = next((x for x in m if x and "Groceries" in [str(v) for v in x]), None)
        ok = hdr is not None
        if ok:
            for row in m:
                if row and isinstance(row[0], str) and row[0] in exp_mon.index:
                    for cat in exp_mon.columns:
                        v = row[hdr.index(cat)]
                        if not isinstance(v, (int, float)) or abs(v - exp_mon.loc[row[0], cat]) > 0.011:
                            ok = False
            ok = ok and all(any(isinstance(r0[0], str) and r0[0] == mm for r0 in m) for mm in exp_mon.index)
        checks["monthly_values"] = ok
    return checks


def check_docx(docx_path):
    """Task C. Returns {check_name: bool}."""
    d = docx.Document(docx_path); checks = {}
    checks["title"] = any(p.style.name == "Title" and p.text.strip() == "Q1 Budget Review" for p in d.paragraphs)
    heads = [p.text.strip() for p in d.paragraphs if p.style.name.startswith("Heading")]
    checks["headings"] = all(h in heads for h in ["Summary", "Budget vs Actual", "Recommendations"])
    checks["bullets_3"] = sum(p.style.name == "List Bullet" for p in d.paragraphs) == 3
    t = d.tables[0] if d.tables else None
    norm = lambda s: re.sub(r"[^\d.A-Za-z ]", "", s).strip()
    want = [["Category", "Budget", "Actual"], ["Rent", "4200", "4200"], ["Groceries", "1500", "1380"], ["Dining", "600", "720"], ["Transport", "900", "600"]]
    checks["table"] = bool(t) and [[norm(c.text) for c in r.cells] for r in t.rows] == want
    checks["header_bold"] = bool(t) and all(any(r.bold for r in c.paragraphs[0].runs) for c in t.rows[0].cells)
    txt = " ".join(p.text for p in d.paragraphs).replace(",", "")
    checks["numbers"] = all(x in txt for x in ["6900", "300", "7200"])
    return checks
