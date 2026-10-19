import docx
d=docx.Document(); d.add_paragraph("Q1 Budget Review",style="Title"); d.add_heading("Summary",1)
d.add_paragraph("Total actual spending of 6,900 was 300 under the total budget of 7,200.")
d.add_heading("Budget vs Actual",1); rows=[("Category","Budget","Actual"),("Rent",4200,4200),("Groceries",1500,1380),("Dining",600,720),("Transport",900,600)]
t=d.add_table(rows=5,cols=3)
for i,r in enumerate(rows):
    for j,v in enumerate(r):
        c=t.cell(i,j); c.text=str(v)
        if i==0: c.paragraphs[0].runs[0].bold=True
d.add_heading("Recommendations",1)
for s in ["Cap dining","Cook more","Track weekly"]: d.add_paragraph(s,style="List Bullet")
d.save("memo.docx")
