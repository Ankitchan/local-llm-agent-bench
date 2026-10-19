from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Title
title = doc.add_heading("Q1 Budget Review", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Summary
doc.add_heading("Summary", 1)
p = doc.add_paragraph("Total actual spending of 6,900 was 300 under the total budget of 7,200. This indicates that spending was well controlled overall, though a few categories warrant attention. The following table breaks down the differences by category.")

# Budget vs Actual
doc.add_heading("Budget vs Actual", 1)
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
hdr_cells = table.rows[0].cells
hdr_cells[0].text = "Category"
hdr_cells[1].text = "Budget"
hdr_cells[2].text = "Actual"
for cell in hdr_cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

data = [
    ("Rent", "4200", "4200"),
    ("Groceries", "1500", "1380"),
    ("Dining", "600", "720"),
    ("Transport", "900", "600"),
]
for category, budget, actual in data:
    row_cells = table.add_row().cells
    row_cells[0].text = category
    row_cells[1].text = budget
    row_cells[2].text = actual

# Recommendations
doc.add_heading("Recommendations", 1)
doc.add_paragraph("Investigate the increase in dining expenses and set a monthly cap to bring it back within budget.", style="List Bullet")
doc.add_paragraph("Encourage meal preparation at home to reduce frequent restaurant spending.", style="List Bullet")
doc.add_paragraph("Review and renegotiate any corporate dining or catering agreements to lower recurring costs.", style="List Bullet")

doc.save("memo.docx")
