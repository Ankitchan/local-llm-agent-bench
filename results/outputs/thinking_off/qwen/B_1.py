import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

# Read the CSV
df = pd.read_csv('transactions.csv', parse_dates=['date'])

# Filter for expenses only (negative amounts)
expenses = df[df['amount'] < 0].copy()
expenses['abs_amount'] = expenses['amount'].abs()

# --- Sheet 1: By Category ---
by_cat = expenses.groupby('category').agg(
    Total_Spent=('abs_amount', 'sum'),
    Transactions=('amount', 'count')
).reset_index()

# Calculate % of Total Spent
total_spent_all = by_cat['Total_Spent'].sum()
by_cat['Pct_of_Total'] = (by_cat['Total_Spent'] / total_spent_all) * 100

# Sort by Total Spent descending
by_cat = by_cat.sort_values(by='Total_Spent', ascending=False)

# Prepare data for writing
cat_headers = ['Category', 'Total Spent', 'Transactions', '% of Total Spent']
cat_data = by_cat.values.tolist()

# --- Sheet 2: Monthly ---
expenses['year_month'] = expenses['date'].dt.to_period('M').astype(str)
monthly_pivot = expenses.pivot_table(
    index='year_month',
    columns='category',
    values='abs_amount',
    aggfunc='sum',
    fill_value=0
)

# Add Total column
monthly_pivot['Total'] = monthly_pivot.sum(axis=1)

# Sort months
monthly_pivot = monthly_pivot.sort_index()

# Create Excel writer
with pd.ExcelWriter('summary.xlsx', engine='openpyxl') as writer:
    # Write Sheet 1: By Category
    by_cat.to_excel(writer, sheet_name='By Category', index=False, header=cat_headers)
    wb = writer.book
    ws1 = wb['By Category']
    
    # Bold header row in Sheet 1
    header_font = Font(bold=True)
    for cell in ws1[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    
    # Format money with 2 decimals in Sheet 1
    for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = '0.00'
    for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row, min_col=4, max_col=4):
        for cell in row:
            cell.number_format = '0.00'
    
    # Add TOTAL row with SUM formula
    total_row_idx = ws1.max_row + 1
    ws1.cell(row=total_row_idx, column=1, value='TOTAL')
    ws1.cell(row=total_row_idx, column=1).font = header_font
    
    # Total Spent cell has SUM formula
    total_spent_cell = ws1.cell(row=total_row_idx, column=2)
    total_spent_cell.value = f'=SUM(B2:B{total_row_idx-1})'
    total_spent_cell.number_format = '0.00'
    total_spent_cell.font = header_font
    
    # Transactions cell has SUM formula
    total_trans_cell = ws1.cell(row=total_row_idx, column=3)
    total_trans_cell.value = f'=SUM(C2:C{total_row_idx-1})'
    total_trans_cell.font = header_font
    
    # % of Total Spent cell is blank or 100%
    total_pct_cell = ws1.cell(row=total_row_idx, column=4)
    total_pct_cell.value = 100
    total_pct_cell.number_format = '0.00'
    total_pct_cell.font = header_font
    
    # Write Sheet 2: Monthly
    monthly_pivot.to_excel(writer, sheet_name='Monthly')
    ws2 = wb['Monthly']
    
    # Bold header row in Sheet 2
    for cell in ws2[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    
    # Format money with 2 decimals in Sheet 2 (all data cells except headers)
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, min_col=1, max_col=ws2.max_column):
        for cell in row:
            if cell.column > 1:  # Skip the month column
                cell.number_format = '0.00'
    
    # Bold the Total column header
    total_col_letter = get_column_letter(ws2.max_column)
    total_header_cell = ws2[f'{total_col_letter}1']
    total_header_cell.font = header_font
