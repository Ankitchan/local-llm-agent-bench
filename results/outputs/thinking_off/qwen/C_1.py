from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Title
title = doc.add_heading('Q1 Budget Review', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Summary
doc.add_heading('Summary', level=1)
summary_text = (
    "Total actual spending was 6,900, which is 300 under the total budget of 7,200. "
    "This indicates that the overall budget was managed effectively during the quarter."
)
doc.add_paragraph(summary_text)

# Budget vs Actual
doc.add_heading('Budget vs Actual', level=1)

table = doc.add_table(rows=1, cols=3, style='Table Grid')
table.style = 'Table Grid'

# Header row
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Category'
hdr_cells[1].text = 'Budget'
hdr_cells[2].text = 'Actual'

# Make header row bold
for cell in hdr_cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.bold = True

# Data rows
data = [
    ('Rent', '4200', '4200'),
    ('Groceries', '1500', '1380'),
    ('Dining', '600', '720'),
    ('Transport', '900', '600'),
]

for category, budget, actual in data:
    row_cells = table.add_row().cells
    row_cells[0].text = category
    row_cells[1].text = budget
    row_cells[2].text = actual

# Recommendations
doc.add_heading('Recommendations', level=1)

recommendations = [
    "Review dining expenses and set a monthly limit to curb overspending.",
    "Encourage meal prepping at home to reduce the frequency of eating out.",
    "Track all dining expenditures weekly to identify and address trends early."
]

for rec in recommendations:
    doc.add_paragraph(rec, style='List Bullet')

doc.save('memo.docx')
