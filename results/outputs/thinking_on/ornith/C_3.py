from docx import Document

doc = Document()

# Title
title = doc.add_heading("Q1 Budget Review", 0)

# Summary
doc.add_heading("Summary", 1)
doc.add_paragraph(
    "Total actual spending of 6,900 came in 300 under the total budget of 7,200. "
    "The quarter was managed within budget, with savings in groceries and transport "
    "offsetting a slight overspend on dining."
)

# Budget vs Actual
doc.add_heading("Budget vs Actual", 1)

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"

header_cells = table.rows[0].cells
header_cells[0].text = "Category"
header_cells[1].text = "Budget"
header_cells[2].text = "Actual"

for cell in header_cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

data_rows = [
    ("Rent", "4200", "4200"),
    ("Groceries", "1500", "1380"),
    ("Dining", "600", "720"),
    ("Transport", "900", "600"),
]

for category, budget, actual in data_rows:
    cells = table.add_row().cells
    cells[0].text = category
    cells[1].text = budget
    cells[2].text = actual

# Recommendations
doc.add_heading("Recommendations", 1)
doc.add_paragraph(
    "Review and reduce dining-out expenses by setting a weekly cap on discretionary meals.",
    style="List Bullet"
)
doc.add_paragraph(
    "Introduce a mid-week no-dining-out policy to encourage home-cooked meals.",
    style="List Bullet"
)
doc.add_paragraph(
    "Negotiate corporate discounts with popular restaurants to lower recurring dining costs.",
    style="List Bullet"
)

doc.save("memo.docx")
