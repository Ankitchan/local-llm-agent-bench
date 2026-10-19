from docx import Document

doc = Document()

# Title
doc.add_paragraph('Q1 Budget Review', style='Title')

# Summary
doc.add_heading('Summary', level=1)
doc.add_paragraph(
    'Total actual spending of 6,900 was 300 under the total budget of 7,200. '
    'This favorable variance demonstrates strong financial discipline, though '
    'targeted adjustments are recommended to sustain performance.'
)

# Budget vs Actual
doc.add_heading('Budget vs Actual', level=1)
table = doc.add_table(rows=5, cols=3)
table.style = 'Table Grid'
headers = ['Category', 'Budget', 'Actual']
for i, header in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = header
    cell.paragraphs[0].runs[0].bold = True

data = [
    ('Rent', '4200', '4200'),
    ('Groceries', '1500', '1380'),
    ('Dining', '600', '720'),
    ('Transport', '900', '600')
]
for i, row_data in enumerate(data):
    for j, text in enumerate(row_data):
        table.rows[i+1].cells[j].text = text

# Recommendations
doc.add_heading('Recommendations', level=1)
recommendations = [
    'Set a strict weekly limit on restaurant meals to curb discretionary spending.',
    'Prepare lunches at home instead of eating out to reduce daily food expenses.',
    'Replace casual dining outings with low-cost social activities to maintain relationships without high costs.'
]
for rec in recommendations:
    doc.add_paragraph(rec, style='List Bullet')

doc.save('memo.docx')
