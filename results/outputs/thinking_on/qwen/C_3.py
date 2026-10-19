from docx import Document

doc = Document()

# Title
doc.add_heading('Q1 Budget Review', level=0)

# Summary
doc.add_heading('Summary', level=1)
doc.add_paragraph(
    "Total actual spending of 6,900 was 300 under the total budget of 7,200. "
    "This favorable variance highlights effective cost control and disciplined expense management."
)

# Budget vs Actual Table
doc.add_heading('Budget vs Actual', level=1)
table = doc.add_table(rows=5, cols=3)
table.style = 'Table Grid'

headers = ['Category', 'Budget', 'Actual']
data = [
    ['Rent', 4200, 4200],
    ['Groceries', 1500, 1380],
    ['Dining', 600, 720],
    ['Transport', 900, 600]
]

# Header row (bold)
for i, header in enumerate(headers):
    cell = table.cell(0, i)
    cell.text = header
    cell.paragraphs[0].runs[0].font.bold = True

# Data rows
for row_idx, row_data in enumerate(data, start=1):
    for col_idx, value in enumerate(row_data):
        table.cell(row_idx, col_idx).text = str(value)

# Recommendations
doc.add_heading('Recommendations', level=1)
recommendations = [
    "Implement a weekly meal prep schedule to reduce reliance on restaurant dining.",
    "Set a strict monthly cap for eating out and track all dining expenses in a budgeting app.",
    "Choose affordable local eateries and limit premium dining to special occasions only."
]
for rec in recommendations:
    doc.add_paragraph(rec, style='List Bullet')

doc.save('memo.docx')
